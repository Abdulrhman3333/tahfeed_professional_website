from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import load_workbook

from .models import Attendance, Role, Student, TeacherAttendance, UserRole


class StudentDataExportTests(TestCase):
    def setUp(self):
        self.manager = User.objects.create_user(username='manager', password='password')
        manager_role, _ = Role.objects.get_or_create(code='manager', defaults={'name': 'المدير'})
        UserRole.objects.create(user=self.manager, role=manager_role)
        self.teacher = User.objects.create_user(
            username='teacher', password='password', first_name='معلم', last_name='مسؤول'
        )
        Student.objects.create(
            full_name='طالب تجريبي',
            student_phone='0501234567',
            parent_phone='0557654321',
            identity_number='1234567890',
            grade='1_pri',
            last_tested_part='1',
            previous_center='مركز سابق',
            neighborhood='حي النور',
            teacher=self.teacher,
            status='منتظم',
        )

    def test_manager_can_download_all_requested_student_fields(self):
        self.client.force_login(self.manager)

        response = self.client.get('/admin-dashboard/export-students/')

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment; filename="student_data.xlsx"', response['Content-Disposition'])
        worksheet = load_workbook(BytesIO(response.content), data_only=False).active
        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            [
                'اسم الطالب',
                'جوال الطالب',
                'جوال ولي الأمر',
                'رقم الهوية الوطنية',
                'الصف الدراسي',
                'آخر جزء تم اختباره',
                'التحفيظ السابق',
                'الحي',
                'المعلم المسؤول',
                'الحالة',
                'المرحلة التعليمية',
            ],
        )
        self.assertEqual(
            [cell.value for cell in worksheet[2]],
            [
                'طالب تجريبي',
                '0501234567',
                '0557654321',
                '1234567890',
                'أول ابتدائي',
                'جزء 1',
                'مركز سابق',
                'حي النور',
                'معلم مسؤول',
                'منتظم',
                'مبكرة',
            ],
        )
        self.assertTrue(worksheet.sheet_view.rightToLeft)

    def test_non_admin_cannot_download_student_data(self):
        teacher = User.objects.get(username='teacher')
        self.client.force_login(teacher)

        response = self.client.get('/admin-dashboard/export-students/')

        self.assertEqual(response.status_code, 302)

    def test_anonymous_user_is_redirected_to_login(self):
        response = self.client.get('/admin-dashboard/export-students/')

        self.assertEqual(response.status_code, 302)
        self.assertIn('/login/', response['Location'])


class AttendanceReportExportTests(TestCase):
    def setUp(self):
        self.manager = User.objects.create_user(username='manager', password='password')
        manager_role, _ = Role.objects.get_or_create(code='manager', defaults={'name': 'المدير'})
        UserRole.objects.create(user=self.manager, role=manager_role)
        self.teacher = User.objects.create_user(
            username='teacher', password='password', first_name='معلم', last_name='مسؤول'
        )
        self.student = Student.objects.create(
            full_name='طالب التقرير',
            student_phone='0500000000',
            parent_phone='0550000000',
            identity_number='1111111111',
            grade='1_pri',
            teacher=self.teacher,
        )
        Attendance.objects.create(
            student=self.student,
            date='2026-01-05',
            weekday='الاثنين',
            week_number=1,
            status='غائب',
        )
        Attendance.objects.create(
            student=self.student,
            date='2026-01-06',
            weekday='الثلاثاء',
            week_number=1,
            status='غياب بعذر',
        )
        Attendance.objects.create(
            student=self.student,
            date='2026-01-07',
            weekday='الأربعاء',
            week_number=1,
            status='متأخر',
        )
        Attendance.objects.create(
            student=self.student,
            date='2026-02-01',
            weekday='الأحد',
            week_number=4,
            status='غائب',
        )
        TeacherAttendance.objects.create(
            teacher=self.teacher,
            date='2026-01-05',
            weekday='الاثنين',
            week_number=1,
            status='غياب بعذر',
        )

    def test_student_report_uses_date_range_and_has_status_columns(self):
        self.client.force_login(self.manager)
        response = self.client.get(
            '/admin-dashboard/export-student-absence/',
            {'start_date': '2026-01-01', 'end_date': '2026-01-31'},
        )

        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(BytesIO(response.content), data_only=False).active
        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            [
                'اسم الطالب', 'عدد أيام الغياب', 'الغياب', 'الغياب بعذر',
                'التأخير', 'جميع الحالات', 'رقم الهوية', 'جوال ولي الأمر',
                'جوال الطالب', 'اسم المعلم', 'الصف', 'المرحلة',
            ],
        )
        self.assertEqual([cell.value for cell in worksheet[2]][:6], [
            'طالب التقرير', 2, 1, 1, 1, 3,
        ])

    def test_teacher_report_has_stage_and_status_columns(self):
        self.client.force_login(self.manager)
        response = self.client.get(
            '/admin-dashboard/export-teacher-absence/',
            {'start_date': '2026-01-01', 'end_date': '2026-01-31'},
        )

        self.assertEqual(response.status_code, 200)
        worksheet = load_workbook(BytesIO(response.content), data_only=False).active
        self.assertEqual([cell.value for cell in worksheet[1]], [
            'اسم المعلم', 'عدد أيام الغياب', 'الغياب', 'الغياب بعذر',
            'التأخير', 'جميع الحالات', 'المرحلة',
        ])
        self.assertEqual([cell.value for cell in worksheet[2]], [
            'معلم مسؤول', 1, 0, 1, 0, 1, 'مبكرة',
        ])

    def test_reports_require_a_valid_date_range(self):
        self.client.force_login(self.manager)

        response = self.client.get('/admin-dashboard/export-student-absence/')

        self.assertEqual(response.status_code, 400)


class PreparerAttendanceAccessTests(TestCase):
    def setUp(self):
        self.supervisor = User.objects.create_user(username='supervisor', password='password')
        self.preparer = User.objects.create_user(username='preparer', password='password')
        supervisor_role, _ = Role.objects.get_or_create(code='supervisor', defaults={'name': 'المشرف'})
        preparer_role, _ = Role.objects.get_or_create(code='preparer', defaults={'name': 'المحضّر'})
        UserRole.objects.create(user=self.supervisor, role=supervisor_role)
        UserRole.objects.create(user=self.preparer, role=preparer_role)

    def test_supervisor_without_preparer_role_is_denied(self):
        self.client.force_login(self.supervisor)

        response = self.client.get('/preparer-attendance/')

        self.assertEqual(response.status_code, 302)

    def test_preparer_can_view_attendance_summary(self):
        self.client.force_login(self.preparer)

        response = self.client.get('/preparer-attendance/')

        self.assertEqual(response.status_code, 200)
